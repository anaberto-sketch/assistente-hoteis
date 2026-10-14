"""Gera data/hotels_cleaned.csv a partir do xlsx do material.

A fonte canonica sao as 20 linhas de hotel da aba "2. Datos y query" do xlsx.
As metricas ja vem pre-computadas do Metabase (query 171721); aqui so
normalizamos a representacao: o xlsx guarda numeros como float e flags como
string ("true"/"false"/None). Escrevemos um CSV com as 59 colunas do field
dictionary, na mesma ordem, descartando a coluna extra `perfil_ejemplo`.

Uso:
    uv run --with openpyxl python3 scripts/regenerate_hotels_csv.py
"""

from __future__ import annotations

import csv
from pathlib import Path

import yaml

REPO_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = REPO_ROOT / "data"
XLSX_PATH = DATA_DIR / "Hackaton AI - Asistente de Revenue para Hoteles (1).xlsx"
DICTIONARY_PATH = DATA_DIR / "field_dictionary.yaml"
CSV_PATH = DATA_DIR / "hotels_cleaned.csv"

SHEET = "2. Datos y query"
DROP_COLUMNS = frozenset({"perfil_ejemplo"})


def _load_dictionary_fields() -> list[str]:
    raw = yaml.safe_load(DICTIONARY_PATH.read_text(encoding="utf-8"))
    return [item["field"] for item in raw]


def _fmt(value: object) -> str:
    """Converte o valor cru do xlsx para a string do CSV."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        # repr e o caminho round-trip mais curto: preserva a precisao exata.
        return repr(value)
    return str(value).strip()


def _read_xlsx_rows():
    import openpyxl

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[SHEET]

    # Linha de cabecalho: primeira em que a primeira celula e "hotel_id".
    header_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "hotel_id":
            header_row = r
            break
    if header_row is None:
        raise RuntimeError("Cabecalho nao encontrado: nenhuma linha comeca com 'hotel_id'.")

    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]

    rows: list[dict[str, object]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        first = ws.cell(r, 1).value
        if first is None:
            # Fim do bloco de dados: abaixo disso vem a query SQL do Metabase.
            break
        if first == "hotel_id":
            # Linha de cabecalho duplicada (o xlsx repete o cabecalho na linha 5).
            continue
        row = {
            headers[c - 1]: ws.cell(r, c).value
            for c in range(1, ws.max_column + 1)
            if headers[c - 1] is not None
        }
        rows.append(row)
    return headers, rows


def main() -> None:
    fields = _load_dictionary_fields()
    headers, rows = _read_xlsx_rows()

    # Colunas no xlsx que nao estao no dicionario.
    extra = [h for h in headers if h is not None and h not in fields]
    dropped = [h for h in extra if h in DROP_COLUMNS]
    unknown = [h for h in extra if h not in DROP_COLUMNS]

    with CSV_PATH.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: _fmt(row.get(field)) for field in fields})

    print(f"Leu {len(rows)} hotéis do xlsx.")
    if dropped:
        print(f"Colunas descartadas: {sorted(dropped)}")
    if unknown:
        print(f"AVISO: colunas no xlsx fora do dicionario (ignoradas): {sorted(unknown)}")
    print(f"Escreveu {CSV_PATH.name} ({len(rows)} linhas x {len(fields)} colunas).")


if __name__ == "__main__":
    main()
